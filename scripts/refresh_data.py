"""Publish reconciled central-budget spending from Ministry of Finance reports."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import time
import urllib.error
import urllib.request
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

FIRST_YEAR = 2018
MINISTRY = "https://www.gov.pl"
REPORT_PAGE = MINISTRY + "/web/finanse/sprawozdania-miesieczne-{year}"
ROMAN_MONTHS = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6,
                "VII": 7, "VIII": 8, "IX": 9, "X": 10, "XI": 11, "XII": 12}
MONTH_HEADER = re.compile(r"^I(?:\s*[-–]\s*(XII|XI|X|IX|VIII|VII|VI|V|IV|III|II))?(?:\s*\*.*)?$", re.I)
CATEGORY_CODES = {
    "grants": "1.", "benefits": "2.", "operations": "3.",
    "capital": "4.", "debtService": "5.", "euContribution": "6.",
    "euProjects": "7.",
}


class Links(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[tuple[str, str]] = []
        self.href: str | None = None
        self.label: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            self.href = dict(attrs).get("href")
            self.label = []

    def handle_data(self, data: str) -> None:
        if self.href is not None:
            self.label.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self.href:
            self.links.append((self.href, " ".join(" ".join(self.label).split())))
            self.href = None


def fetch(url: str) -> bytes:
    cache = Path(__file__).resolve().parents[1] / ".context/source-cache"
    cache.mkdir(parents=True, exist_ok=True)
    cache_path = cache / hashlib.sha256(url.encode()).hexdigest()
    if "/attachment/" in url and cache_path.exists():
        return cache_path.read_bytes()
    request = urllib.request.Request(url, headers={"User-Agent": "Polstat/1.0 (+public data attribution)"})
    for attempt in range(3):
        try:
            with urllib.request.urlopen(request, timeout=180) as response:
                data = response.read()
            break
        except TimeoutError:
            if attempt == 2:
                raise
            time.sleep(2 ** attempt)
    if not data:
        raise ValueError(f"Empty source: {url}")
    if "/attachment/" in url:
        cache_path.write_bytes(data)
    return data


def source_link(page: str) -> tuple[str, str | None]:
    parser = Links()
    parser.feed(fetch(page).decode("utf-8"))
    for href, label in parser.links:
        if "Sprawozdanie operatywne" in label and ".xlsx" in label.lower():
            match = re.search(r"(?<!\d)(20\d{6})(?!\d)", label)
            source_file_date = None
            if match:
                try:
                    source_file_date = datetime.strptime(match.group(1), "%Y%m%d").date().isoformat()
                except ValueError:
                    pass
            return urllib.request.urljoin(MINISTRY, href), source_file_date
    raise ValueError(f"No operational XLSX report on {page}")


def sheet(workbook, title: str):
    return workbook[next(name for name in workbook.sheetnames if " ".join(name.split()) == title)]


def number(value: object, scale: float) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Expected a numeric source cell, got {value!r}")
    return round(float(value) / scale, 6)


def blocks(worksheet) -> dict[int, tuple[list[tuple[object, ...]], int, int, int]]:
    rows = list(worksheet.iter_rows(values_only=True))
    headers: list[tuple[int, int, int]] = []
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            if col_index < 2 or not isinstance(value, str):
                continue
            match = MONTH_HEADER.fullmatch(value.strip())
            if match:
                month = ROMAN_MONTHS[match.group(1).upper()] if match.group(1) else 1
                headers.append((month, row_index, col_index))
    if not headers:
        raise ValueError(f"No monthly columns in {worksheet.title}")
    result = {}
    for month, row_index, col_index in headers:
        next_header = min((r for _, r, _ in headers if r > row_index + 2), default=len(rows))
        result[month] = (rows, row_index + 1, min(next_header, row_index + 50), col_index)
    return result


def value_for(block, predicate, scale: float) -> float:
    rows, start, end, col = block
    for row in rows[start:end]:
        if predicate(row):
            return number(row[col], scale)
    raise ValueError(f"Missing spending cell near source row {start + 1}")


def parse_functions(worksheet, expected_total: float, period: str, url: str,
                    source_file_date: str | None) -> dict:
    rows = list(worksheet.iter_rows(values_only=True))
    overall = next(
        i for i, row in enumerate(rows)
        if len(row) > 4 and row[3] == "c" and isinstance(row[4], (int, float))
    )
    raw_total = float(rows[overall][4])
    # Historical sheets sometimes store displayed thousands and sometimes PLN.
    # Reconcile against the independently parsed monthly total before scaling.
    scale = next((candidate for candidate in (1_000, 1_000_000)
                  if abs(raw_total / candidate - expected_total) < 0.01), None)
    if scale is None:
        raise ValueError(f"Unrecognised functional-spending scale in {period}: {raw_total}")
    functions = []
    for index, row in enumerate(rows):
        code = row[0] if row else None
        label = row[2] if len(row) > 2 else None
        if not (isinstance(code, str) and len(code.strip()) == 3 and
                code.strip().isdigit() and isinstance(label, str) and label.strip()):
            continue
        measures = {candidate[3]: (0.0 if candidate[4] is None or
                                  isinstance(candidate[4], str) and not candidate[4].strip()
                                  else number(candidate[4], scale))
                    for candidate in rows[index:index + 5]
                    if len(candidate) > 4 and candidate[3] in ("a", "b", "c")}
        if set(measures) != {"a", "b", "c"}:
            raise ValueError(f"Incomplete functional row {code} in {period}")
        functions.append({
            "code": code.strip(), "label": " ".join(label.split()),
            "actual": measures["c"], "originalPlan": measures["a"],
            "amendedPlan": measures["b"],
        })
    if not functions or len({item["code"] for item in functions}) != len(functions):
        raise ValueError(f"Missing or duplicate spending functions in {period}")
    difference = round(expected_total - sum(item["actual"] for item in functions), 6)
    if abs(difference) > 2:
        raise ValueError(f"Functional spending does not reconcile in {period}: {difference}")
    result = {
        "period": period, "total": expected_total,
        "originalPlan": number(rows[overall - 2][4], scale),
        "amendedPlan": number(rows[overall - 1][4], scale),
        "functions": functions, "unit": "PLN million", "sourceUrl": url,
        "publishedAt": None, "sourceFileDate": source_file_date,
        "status": "reported_actual",
    }
    if abs(difference) > 0.01:
        result["sourceDifference"] = difference
    return result


def parse_report(data: bytes, year: int, url: str,
                 source_file_date: str | None) -> tuple[list[dict], dict]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    general = blocks(sheet(workbook, "TABLICA 1"))
    economic = blocks(sheet(workbook, "TABLICA 6"))
    months = sorted(
        month for month in set(general) & set(economic)
        if any(isinstance(row[0], str) and
               re.match(r"^II\.\s*WYDATKI", row[0].strip(), re.I) and
               isinstance(row[general[month][3]], (int, float))
               for row in general[month][0][general[month][1]:general[month][2]])
    )
    if not months or months != list(range(1, max(months) + 1)):
        raise ValueError(f"Spending month gap in {year}: {months}")
    spending = []
    for month in months:
        g, e = general[month], economic[month]
        total = value_for(
            g, lambda row: isinstance(row[0], str)
            and bool(re.match(r"^II\.\s*WYDATKI", row[0].strip(), re.I)), 1_000
        )
        economic_total = value_for(
            e, lambda row: isinstance(row[1], str) and
            "WYDATKI OGÓŁEM" in row[1], 1_000_000
        )
        if abs(total - economic_total) > 0.003:
            raise ValueError(f"Spending totals differ in {year}-{month:02d}")
        categories = {
            name: value_for(e, lambda row, code=code: row[1] == code, 1_000_000)
            for name, code in CATEGORY_CODES.items()
        }
        difference = round(total - sum(categories.values()), 6)
        if abs(difference) > 0.003:
            if abs(difference) > 2:
                raise ValueError(f"Spending categories do not reconcile in {year}-{month:02d}")
            categories["sourceDifference"] = difference
        spending.append({
            "period": f"{year}-{month:02d}", "total": total,
            "categories": categories, "unit": "PLN million", "sourceUrl": url,
            "publishedAt": None, "sourceFileDate": source_file_date,
            "status": "reported_actual",
        })
    last = spending[-1]
    functions = parse_functions(sheet(workbook, "TABLICA 7"), last["total"],
                                last["period"], url, source_file_date)
    workbook.close()
    return spending, functions


def build_snapshot(today: date | None = None) -> dict:
    today = today or date.today()
    spending, breakdowns = [], []
    for year in range(FIRST_YEAR, today.year + 1):
        try:
            url, source_file_date = source_link(REPORT_PAGE.format(year=year))
        except (ValueError, urllib.error.HTTPError):
            if year == today.year and today.month <= 3 and spending and spending[-1]["period"] == f"{year - 1}-12":
                break
            raise
        records, functions = parse_report(fetch(url), year, url, source_file_date)
        if year < today.year and records[-1]["period"] != f"{year}-12":
            raise ValueError(f"Historical year {year} lacks December")
        spending.extend(records)
        breakdowns.append(functions)
    if today.month > 3 and spending[-1]["period"] < f"{today.year}-01":
        raise ValueError("Current-year spending data is missing")
    return {
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "unit": "PLN million", "spending": spending, "breakdowns": breakdowns,
    }


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_snapshot(snapshot: dict, root: Path) -> None:
    path = root / "src/data/snapshot.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        current = json.loads(path.read_text())
        comparable = lambda data: {key: value for key, value in data.items() if key != "generatedAt"}
        if comparable(current) == comparable(snapshot):
            snapshot = current
        else:
            path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n")
    else:
        path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n")
    data_dir = root / "public/data"
    data_dir.mkdir(parents=True, exist_ok=True)
    category_fields = list(CATEGORY_CODES) + ["sourceDifference"]
    cumulative = [{**row, **row["categories"]} for row in snapshot["spending"]]
    source_fields = ["unit", "sourceUrl", "publishedAt", "sourceFileDate", "status"]
    write_csv(data_dir / "wydatki-narastajaco.csv",
              ["period", "total", *category_fields, *source_fields], cumulative)
    monthly = []
    previous = None
    for row in snapshot["spending"]:
        if previous and previous["period"][:4] != row["period"][:4]:
            previous = None
        monthly.append({
            "period": row["period"],
            "total": round(row["total"] - (previous["total"] if previous else 0), 6),
            "unit": row["unit"], "sourceUrl": row["sourceUrl"],
            "publishedAt": row["publishedAt"],
            "sourceFileDate": row["sourceFileDate"], "status": row["status"],
        })
        previous = row
    write_csv(data_dir / "wydatki-miesiecznie.csv", ["period", "total", *source_fields], monthly)
    functions = [
        {"period": group["period"], **item, "unit": group["unit"],
         "sourceUrl": group["sourceUrl"], "publishedAt": group["publishedAt"],
         "sourceFileDate": group["sourceFileDate"],
         "status": group["status"]}
        for group in snapshot["breakdowns"] for item in group["functions"]
    ]
    write_csv(data_dir / "wydatki-dzialy.csv",
              ["period", "code", "label", "actual", "originalPlan", "amendedPlan", *source_fields],
              functions)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    snapshot = build_snapshot()
    write_snapshot(snapshot, args.root)
    print(f"Imported {len(snapshot['spending'])} spending months and {len(snapshot['breakdowns'])} functional breakdowns")
    print(f"Latest period: {snapshot['spending'][-1]['period']}")
